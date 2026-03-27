from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
import pandas as pd
import time
import os
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

def preencher_formulario():
    # Caminho absoluto para o Excel
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    excel_path = os.path.join(base_dir, 'rpa', 'dados', 'faturas.xlsx')

    if not os.path.exists(excel_path):
        print(f"Erro: Arquivo {excel_path} não encontrado.")
        return

    options = webdriver.ChromeOptions()
    options.add_argument('--headless')
    options.add_argument('--no-sandbox')
    options.add_argument('--disable-dev-shm-usage')

    try:
        with webdriver.Chrome(options=options) as driver:
            driver.get("http://127.0.0.1:8000")
            wait = WebDriverWait(driver, 10)

            df = pd.read_excel(excel_path)
            wb = load_workbook(excel_path)
            ws = wb.active

            green_fill = PatternFill(start_color='00FF00', end_color='00FF00', fill_type='solid')

            for index, row in df.iterrows():
                try:
                    nome_field = wait.until(EC.element_to_be_clickable((By.NAME, "nome")))
                    nome_field.clear()
                    nome_field.send_keys(str(row.get('nome', 'Cliente '+str(index))))

                    cpf_field = driver.find_element(By.NAME, "cpf")
                    cpf_field.clear()
                    cpf_field.send_keys(str(row['cpf']))

                    email_field = driver.find_element(By.NAME, "email")
                    email_field.clear()
                    email_field.send_keys(str(row.get('email', 'email@exemplo.com')))

                    telefone_field = driver.find_element(By.NAME, "telefone")
                    telefone_field.clear()
                    telefone = str(row.get('telefone', '0000000000'))
                    telefone = "".join(filter(str.isdigit, telefone))
                    telefone_field.send_keys(telefone)

                    endereco_field = driver.find_element(By.NAME, "endereco")
                    endereco_field.clear()
                    endereco_field.send_keys(str(row.get('endereco', 'Endereço não informado')))

                    submit_btn = driver.find_element(By.XPATH, "//button[@value='cadastrar']")
                    submit_btn.click()

                    wait.until(EC.presence_of_element_located((By.CSS_SELECTOR, ".sucesso, .erro")))

                    page = driver.page_source.lower()
                    if "sucesso" in page or "já cadastrado" in page:
                        print(f"Linha {index+2}: Processado com sucesso.")
                        for cell in ws[index + 2]:
                            cell.fill = green_fill
                    else:
                        print(f"Linha {index+2}: Erro detectado na página.")

                except Exception as e:
                    print(f"Erro ao processar linha {index+2}: {e}")

            wb.save(excel_path)
            print("Processo concluído e Excel atualizado.")

    except Exception as e:
        print(f"Erro ao iniciar Selenium: {e}")

if __name__ == "__main__":
    preencher_formulario()
